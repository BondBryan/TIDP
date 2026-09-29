"""Extract the TIDP Library tab and Value lists from the 2026 xlsm into JSON for the web app."""
import json, re, sys, datetime
import openpyxl

SRC = sys.argv[1] if len(sys.argv) > 1 else "BB_TaskInformationDeliveryPlan-2026.xlsm"
wb = openpyxl.load_workbook(SRC, keep_vba=True)
lib = wb["Library"]
vl = wb["Value lists"]

hdr = [c.value for c in lib[1]]
col = {h: i for i, h in enumerate(hdr) if h}
SECTORS = hdr[hdr.index("Education"): hdr.index("FloorDup")]
STAGES = ["Preparation & Brief", "Concept Design", "Developed Design", "Technical Design",
          "Construction", "Handover & Close Out", "In Use"]

def txt(v):
    if v is None: return ""
    if isinstance(v, bool): return "TRUE" if v else "FALSE"
    if hasattr(v, "text"): return ""  # array formula (Role) - computed in app
    return str(v).strip()

def flag(v):
    return txt(v).upper() != "FALSE"   # blank counts as TRUE (macro convention)

EXISTING_RX = re.compile(r"existing|demolition|survey|as[- ]built|as existing", re.I)

docs = []
for r in lib.iter_rows(min_row=2, values_only=True):
    if not txt(r[col["Tab"]]): continue
    tab = txt(r[col["Tab"]])
    number = txt(r[col["Number"]])
    if number.isdigit(): number = number.zfill(4)
    desc = txt(r[col["Description"]])
    wp = txt(r[col["Work Package"]])
    stages = {}
    for i, s in enumerate(STAGES, start=1):
        v = txt(r[col[s]])
        stages[str(i)] = v if v else None
    d = {
        "id": f"{tab}-{txt(r[col['Type']])}-{number}",
        "tab": tab,
        "series": txt(r[col["Series"]]),
        "originator": txt(r[col["Originator"]]) or "BBA",
        "volume": txt(r[col["Volume / system"]]) or "ZZ",
        "level": txt(r[col["Level"]]) or "ZZ",
        "type": txt(r[col["Type"]]),
        "number": number,
        "description": desc,
        "format": txt(r[col["Format"]]),
        "scale": txt(r[col["Scale"]]),
        "workPackage": wp,
        "stages": stages,
        "tender": {"3+": bool(txt(r[col["3+"]])), "4a": bool(txt(r[col["4a"]])), "4b": bool(txt(r[col["4b"]]))},
        "status": txt(r[col["Status"]]),
        "comments": txt(r[col["Outsourcing Comments"]]),
        "sectors": {s: flag(r[col[s]]) for s in SECTORS},
        "floorDup": txt(r[col["FloorDup"]]).upper() == "TRUE",
        "volDup": txt(r[col["VolDup"]]).upper() == "TRUE",
        # NEW field for the web app: "any" | "existing" (only when the project has existing buildings / demolition) | "new" (new build only)
        "buildType": "existing" if EXISTING_RX.search(desc + " " + wp) else "any",
    }
    docs.append(d)

def column(ws, letter, start=3):
    out = []
    for c in ws[letter][start-1:]:
        if c.value is None: continue
        v = c.value
        if isinstance(v, str) and v.startswith("="): continue
        out.append(str(v).strip())
    return out

lists = {
    "sectors": SECTORS,
    "levels": column(vl, "D"),
    "volumes": column(vl, "C"),
    "types": column(vl, "E"),
    "roles": column(vl, "F"),
    "workPackages": column(vl, "H"),
    "formats": column(vl, "I"),
    "statuses": column(vl, "L"),
    "stages": [{"n": i + 1, "name": s} for i, s in enumerate(STAGES)],
    "storeyNames": {"B3": "Basement 3", "B2": "Basement 2", "B1": "Basement 1", "00": "Ground Floor",
                    "01": "First Floor", "02": "Second Floor", "03": "Third Floor", "04": "Fourth Floor",
                    "05": "Fifth Floor", "06": "Sixth Floor", "RF": "Roof", "R2": "Roof 2", "R3": "Roof 3"},
    "tabs": ["Drawings", "Images", "Lists", "Models", "Text", "Video", "Internal"],
    "roleBySector": {"Landscape": "L", "Interior Design": "I", "default": "A"},
    # discipline name used in the exported file names (..._Architect.xlsx)
    "disciplineBySector": {"Landscape": "Landscape", "Interior Design": "Interior Design", "default": "Architect"},
}

meta = {
    "name": "Bond Bryan TIDP Library",
    "version": "2026.2",
    "source": SRC,
    "extracted": datetime.date.today().isoformat(),
    "documents": len(docs),
    "note": "Master list of Task Information Delivery Plan deliverables. Edit in GitHub; the web app reads this file at load.",
}
json.dump({"meta": meta, "lists": lists, "documents": docs}, open("library.json", "w"), indent=1, ensure_ascii=False)
print(len(docs), "documents;", sum(d["buildType"] == "existing" for d in docs), "flagged existing-only")
for d in docs:
    if d["buildType"] == "existing": print("  ", d["tab"], d["number"], d["description"], "|", d["workPackage"])
from collections import Counter
print(Counter(d["tab"] for d in docs))
print(Counter(d["status"] for d in docs))
